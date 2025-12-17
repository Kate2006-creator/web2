from django.http import HttpResponse
from docx import Document
from datetime import datetime
import io

class WordExporter:

    def export_to_word(data, title, headers, fields):
        doc = Document()
        doc.add_heading(title, 0)
        
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        
        for i, header in enumerate(headers): #enumerate - создает пары (индекс, значение):
            #(0, 'Названия'), (1, 'ФИО') и тд
            table.rows[0].cells[i].text = header
        
        for item in data:
            row = table.add_row().cells
            for i, field in enumerate(fields): #проходимся по парам
                row[i].text = str(getattr(item, field, '')) #вытаскиваем наименование ЧЕРЕЗ STR не работает ОШИБКА
        
        return doc
    

    def create_http_response(doc, filename):
        file_stream = io.BytesIO()
        doc.save(file_stream)
        
        response = HttpResponse(
            file_stream.getvalue(),
            content_type='application/octet-stream'  #скачивание
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response